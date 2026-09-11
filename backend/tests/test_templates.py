import unittest
from backend.services.excel_template_service import ExcelTemplateService
from backend.ocr.base import OCRResult, OCRLine
from backend.extractors.dni_extractor import extract_dni_data
from backend.extractors.validators import normalize_civil_status, normalize_and_validate_date


class TestExcelTemplatesAndDni(unittest.TestCase):
    def test_civil_status_normalization(self):
        valid, val, _ = normalize_civil_status("S")
        self.assertTrue(valid)
        self.assertEqual(val, "SOLTERO")

        valid, val, _ = normalize_civil_status("C")
        self.assertTrue(valid)
        self.assertEqual(val, "CASADO")

        valid, val, _ = normalize_civil_status("SOLTERO")
        self.assertTrue(valid)
        self.assertEqual(val, "SOLTERO")

    def test_date_normalization_space_and_compact(self):
        valid, val, _ = normalize_and_validate_date("02 061997")
        self.assertTrue(valid)
        self.assertEqual(val, "02/06/1997")

        valid, val, _ = normalize_and_validate_date("20092023")
        self.assertTrue(valid)
        self.assertEqual(val, "20/09/2023")

        valid, val, _ = normalize_and_validate_date("20 09 2031")
        self.assertTrue(valid)
        self.assertEqual(val, "20/09/2031")

    def test_dni_multiline_betty_valeriana(self):
        front_lines = [
            OCRLine(text="DNI 71795391", confidence=95.0),
            OCRLine(text="DOCUMENTO NACIONAL DE IDENTIDAD", confidence=92.0),
            OCRLine(text="Primer Apellido", confidence=90.0),
            OCRLine(text="FechaInscripcion", confidence=88.0),
            OCRLine(text="MENDEZ", confidence=92.0),
            OCRLine(text="08 05 2008", confidence=89.0),
            OCRLine(text="Fecha Emision", confidence=91.0),
            OCRLine(text="Segundo Apellido", confidence=90.0),
            OCRLine(text="20092023", confidence=93.0),
            OCRLine(text="LOPEZ", confidence=91.0),
            OCRLine(text="FechaCaducidad", confidence=89.0),
            OCRLine(text="20 09 2031", confidence=94.0),
            OCRLine(text="Pre Nombres", confidence=90.0),
            OCRLine(text="BETTY VALERIANA", confidence=92.0),
            OCRLine(text="Nacimiento. Fecha y Ubigeo", confidence=88.0),
            OCRLine(text="02 061997", confidence=90.0),
            OCRLine(text="021505", confidence=88.0),
            OCRLine(text="Estado Civil: S", confidence=90.0),
            OCRLine(text="Sexo: F", confidence=95.0),
        ]
        back_lines = [
            OCRLine(text="CONSTANCIA DE SUFRAGIO", confidence=90.0),
            OCRLine(text="Provincla", confidence=85.0),
            OCRLine(text="Departamento", confidence=88.0),
            OCRLine(text="QUILLO", confidence=89.0),
            OCRLine(text="YUNGAY", confidence=91.0),
            OCRLine(text="ANCASH", confidence=93.0),
            OCRLine(text="Direccion", confidence=90.0),
            OCRLine(text="CASERIO CORACOLLO SN", confidence=92.0),
            OCRLine(text="CuartoNivel 021505", confidence=89.0),
            OCRLine(text="Observaciones", confidence=85.0),
            OCRLine(text="DonaciondeOrganos NO", confidence=87.0),
        ]
        front_ocr = OCRResult(lines=front_lines, full_text="\n".join([l.text for l in front_lines]))
        back_ocr = OCRResult(lines=back_lines, full_text="\n".join([l.text for l in back_lines]))

        data, fields, warnings = extract_dni_data(front_ocr, back_ocr)

        self.assertEqual(data["doc_number"], "71795391")
        self.assertEqual(data["paternal_surname"], "MENDEZ")
        self.assertEqual(data["first_names"], "BETTY VALERIANA")
        self.assertEqual(data["birth_date"], "02/06/1997")
        self.assertEqual(data["issue_date"], "20/09/2023")
        self.assertEqual(data["expiry_date"], "20/09/2031")
        self.assertEqual(data["civil_status"], "SOLTERO")
        self.assertEqual(data["sex"], "F")
        self.assertEqual(data["address"], "CASERIO CORACOLLO SN")
        self.assertIn("021505", data["ubigeo"])
        self.assertTrue(fields["issue_date"].is_valid)
        self.assertTrue(fields["civil_status"].is_valid)

    def test_excel_template_service_lifecycle(self):
        svc = ExcelTemplateService()
        templates = svc.list_templates()
        self.assertGreaterEqual(len(templates), 1)

        tpl_names = [t["filename"] for t in templates]
        target_template = tpl_names[0]

        # Test Batch Fill
        sample_dni = {
            "doc_number": "71795391",
            "paternal_surname": "MENDEZ",
            "maternal_surname": "LOPEZ",
            "first_names": "BETTY VALERIANA",
            "birth_date": "02/06/1997",
            "sex": "F",
            "civil_status": "SOLTERO",
            "address": "CASERIO CORACOLLO SN",
            "department": "ANCASH",
            "province": "YUNGAY",
            "district": "QUILLO",
            "ubigeo": "ANCASH / YUNGAY / QUILLO (021505)",
            "issue_date": "20/09/2023",
            "expiry_date": "20/09/2031",
            "nationality": "PERUANA",
        }

        out_fn, out_path, num_workers, yellow_cols = svc.fill_template_batch(target_template, [sample_dni])
        self.assertTrue(out_path.exists())
        self.assertTrue(out_fn.endswith(".xlsx"))
        self.assertEqual(num_workers, 1)

    def test_excel_template_crud_operations(self):
        svc = ExcelTemplateService()
        
        # 1. Create a dummy test template
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "TEST_SHEET"
        ws.append(["ID", "DOC_PERSONA", "NOMBRE_TRABAJADOR", "ESTADO"])
        dummy_path = svc.templates_dir / "TEST_CRUD_TEMPLATE.xlsx"
        wb.save(dummy_path)
        wb.close()

        # 2. Analyze
        analysis = svc.analyze_template("TEST_CRUD_TEMPLATE.xlsx")
        self.assertEqual(analysis["filename"], "TEST_CRUD_TEMPLATE.xlsx")
        self.assertEqual(len(analysis["sheets"]), 1)

        # 3. Save Custom Mapping
        updated = svc.save_template_mapping("TEST_CRUD_TEMPLATE.xlsx", "TEST_SHEET", {
            "2": "doc_number",
            "3": "first_names",
            "4": "civil_status"
        })
        self.assertEqual(updated["sheets"][0]["mapped_fields"]["doc_number"]["column_index"], 2)

        # 4. Rename
        renamed = svc.rename_template("TEST_CRUD_TEMPLATE.xlsx", "TEST_CRUD_RENAMED.xlsx")
        self.assertEqual(renamed["filename"], "TEST_CRUD_RENAMED.xlsx")
        self.assertTrue((svc.templates_dir / "TEST_CRUD_RENAMED.xlsx").exists())
        self.assertFalse(dummy_path.exists())

        # 5. Fill with custom mappings
        sample_dni = {
            "doc_number": "12345678",
            "first_names": "JUAN CARLOS",
            "civil_status": "CASADO",
        }
        out_fn, out_path = svc.fill_template("TEST_CRUD_RENAMED.xlsx", sample_dni, sheet_name="TEST_SHEET")
        self.assertTrue(out_path.exists())

        # 6. Delete
        deleted = svc.delete_template("TEST_CRUD_RENAMED.xlsx")
        self.assertTrue(deleted)
        self.assertFalse((svc.templates_dir / "TEST_CRUD_RENAMED.xlsx").exists())

    def test_yellow_cells_batch_filling_rule(self):
        """Verifies that only yellow cells receive DNI data and non-yellow columns retain fixed template values."""
        import openpyxl
        from openpyxl.styles import PatternFill
        svc = ExcelTemplateService()

        # Create custom template with Yellow and Non-Yellow columns
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "YELLOW_TEST"
        ws.append(["ITEM", "DNI", "TRABAJADOR", "CODIGO_FIJO", "CARGO_FIJO", "BONO_VARIABLE"])
        
        # Col 2 (DNI), Col 3 (TRABAJADOR), Col 6 (BONO_VARIABLE) are YELLOW
        yellow_fill = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
        ws.cell(row=1, column=2).fill = yellow_fill
        ws.cell(row=1, column=3).fill = yellow_fill
        ws.cell(row=1, column=6).fill = yellow_fill

        # Baseline sample data row with fixed codes
        ws.append([1, "", "", "COD_PLAN_01", "OPERARIO", ""])
        
        tpl_name = "TEST_YELLOW_RULE_TPL.xlsx"
        tpl_path = svc.templates_dir / tpl_name
        wb.save(tpl_path)
        wb.close()

        try:
            workers = [
                {"doc_number": "11111111", "full_name": "TRABAJADOR UNO"},
                {"doc_number": "22222222", "full_name": "TRABAJADOR DOS"}
            ]

            out_fn, out_path, num_workers, yellow_cols = svc.fill_template_batch(tpl_name, workers)
            self.assertTrue(out_path.exists())
            self.assertEqual(num_workers, 2)

            # Inspect generated workbook
            res_wb = openpyxl.load_workbook(out_path)
            res_ws = res_wb["YELLOW_TEST"]

            # Row 2 (Worker 1)
            self.assertEqual(res_ws.cell(row=2, column=1).value, 1) # Item
            self.assertEqual(res_ws.cell(row=2, column=2).value, "11111111") # DNI (Yellow)
            self.assertEqual(res_ws.cell(row=2, column=3).value, "TRABAJADOR UNO") # Full name (Yellow)
            self.assertEqual(res_ws.cell(row=2, column=4).value, "COD_PLAN_01") # Non-yellow fixed preserved!
            self.assertEqual(res_ws.cell(row=2, column=5).value, "OPERARIO") # Non-yellow fixed preserved!
            self.assertIn(res_ws.cell(row=2, column=6).value, [None, ""]) # Yellow variable without info left BLANK!

            # Row 3 (Worker 2)
            self.assertEqual(res_ws.cell(row=3, column=1).value, 2) # Item
            self.assertEqual(res_ws.cell(row=3, column=2).value, "22222222") # DNI (Yellow)
            self.assertEqual(res_ws.cell(row=3, column=3).value, "TRABAJADOR DOS") # Full name (Yellow)
            self.assertEqual(res_ws.cell(row=3, column=4).value, "COD_PLAN_01") # Non-yellow fixed preserved!
            self.assertEqual(res_ws.cell(row=3, column=5).value, "OPERARIO") # Non-yellow fixed preserved!
            self.assertIn(res_ws.cell(row=3, column=6).value, [None, ""]) # Yellow variable without info left BLANK!

            res_wb.close()
        finally:
            if tpl_path.exists():
                tpl_path.unlink()


if __name__ == "__main__":
    unittest.main()

