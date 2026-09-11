import unittest
from pathlib import Path
import json
import tempfile
import openpyxl
from openpyxl.styles import PatternFill

from backend.services.supabase_template_service import SupabaseTemplateService, supabase_template_service
from backend.services.excel_template_service import ExcelTemplateService


class TestSupabaseTemplateService(unittest.TestCase):

    def setUp(self):
        self.service = SupabaseTemplateService()
        self.temp_dir = tempfile.TemporaryDirectory()
        self.temp_path = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_is_configured_and_config_loading(self):
        # When unconfigured
        self.service.supabase_url = ""
        self.service.supabase_key = ""
        self.assertFalse(self.service.is_configured())

        # When configured
        self.service.supabase_url = "https://xyz.supabase.co"
        self.service.supabase_key = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9..."
        self.assertTrue(self.service.is_configured())

    def test_inspect_excel_file_detects_yellow_cells(self):
        # Create a sample workbook with yellow cells
        test_file = self.temp_path / "Test_Template.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "CARGA_DATOS"

        # Headers
        ws.append(["ID", "DNI", "PATERNO", "MATERNO", "NOMBRES", "FECHA_NAC", "SUELDO_BASE"])

        # Row 2 with yellow cells in DNI (col 2), PATERNO (col 3), MATERNO (col 4), NOMBRES (col 5)
        yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        ws.append([1, "", "", "", "", "1990-01-01", 1500])

        ws.cell(row=2, column=2).fill = yellow_fill
        ws.cell(row=2, column=3).fill = yellow_fill
        ws.cell(row=2, column=4).fill = yellow_fill
        ws.cell(row=2, column=5).fill = yellow_fill

        wb.save(str(test_file))

        metadata = self.service.inspect_excel_file(test_file)
        self.assertEqual(len(metadata), 1)
        sheet_info = metadata[0]
        self.assertEqual(sheet_info["sheet_name"], "CARGA_DATOS")
        self.assertEqual(sheet_info["total_columns"], 7)

        yellow_col_indices = [c["col_index"] for c in sheet_info["yellow_columns"]]
        self.assertIn(2, yellow_col_indices)
        self.assertIn(3, yellow_col_indices)
        self.assertIn(4, yellow_col_indices)
        self.assertIn(5, yellow_col_indices)
        self.assertEqual(sheet_info["yellow_columns_count"], 4)

        # Mapped fields detection
        self.assertIn("doc_number", sheet_info["mapped_fields"])
        self.assertIn("paternal_surname", sheet_info["mapped_fields"])
        self.assertIn("maternal_surname", sheet_info["mapped_fields"])
        self.assertIn("first_names", sheet_info["mapped_fields"])

    def test_fill_template_batch_from_file_preserves_formulas(self):
        # Create template with formula in column 4 (=A2*2)
        test_file = self.temp_path / "Batch_Template.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DATOS"
        ws.append(["ITEM", "DNI", "NOMBRE_COMPLETO", "DUPLICADO"])
        ws.append([1, "", "", "=A2*2"])

        yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        ws.cell(row=2, column=2).fill = yellow_fill
        ws.cell(row=2, column=3).fill = yellow_fill
        wb.save(str(test_file))

        # Sample workers
        workers = [
            {
                "doc_number": "42694271",
                "paternal_surname": "GARCIA",
                "maternal_surname": "BRANDAN",
                "first_names": "NADIA CELESTE",
                "full_name": "GARCIA BRANDAN NADIA CELESTE",
            },
            {
                "doc_number": "12345678",
                "paternal_surname": "PEREZ",
                "maternal_surname": "LOPEZ",
                "first_names": "JUAN",
                "full_name": "PEREZ LOPEZ JUAN",
            },
        ]

        excel_service = ExcelTemplateService()
        out_filename, out_path, num_workers, yellow_headers = excel_service.fill_template_batch_from_file(
            test_file,
            "Batch_Template.xlsx",
            workers,
        )

        self.assertTrue(out_path.exists())
        self.assertEqual(num_workers, 2)
        self.assertIn("Batch_Template", out_filename)

        # Inspect generated workbook
        filled_wb = openpyxl.load_workbook(str(out_path), data_only=False)
        filled_ws = filled_wb["DATOS"]

        # Row 2 (first worker)
        self.assertEqual(filled_ws.cell(row=2, column=2).value, "42694271")
        # Row 3 (second worker)
        self.assertEqual(filled_ws.cell(row=3, column=2).value, "12345678")
        # Formula preservation in Row 2
        self.assertEqual(filled_ws.cell(row=2, column=4).value, "=A2*2")


if __name__ == "__main__":
    unittest.main()
