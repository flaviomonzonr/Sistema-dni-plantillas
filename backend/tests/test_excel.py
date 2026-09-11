import unittest
import openpyxl
from pathlib import Path
from backend.storage.excel_manager import append_record_to_excel, generate_filtered_excel
from backend.config import DATA_DIR


class TestExcelManager(unittest.TestCase):
    def setUp(self):
        self.test_excel = DATA_DIR / "test_registros.xlsx"
        if self.test_excel.exists():
            self.test_excel.unlink()

    def tearDown(self):
        if self.test_excel.exists():
            self.test_excel.unlink()

    def test_incremental_append(self):
        rec1 = {
            "doc_type": "DNI",
            "paternal_surname": "GARCIA",
            "maternal_surname": "LOPEZ",
            "first_names": "CARLOS",
            "doc_number": "12345678",
            "nationality": "PERUANA",
            "birth_date": "10/10/1990",
            "sex": "M",
            "address": "AV. LARCO 100",
            "issue_date": "01/01/2020",
            "expiry_date": "01/01/2028",
            "observations": "",
        }
        rec2 = {
            "doc_type": "Carnet de Extranjería",
            "paternal_surname": "SILVA",
            "maternal_surname": "",
            "first_names": "ANA",
            "doc_number": "009876543",
            "nationality": "COLOMBIANA",
            "birth_date": "15/05/1995",
            "sex": "F",
            "address": "CALLE 2 SAN ISIDRO",
            "issue_date": "02/02/2021",
            "expiry_date": "02/02/2026",
            "observations": "Falta apellido materno",
        }

        append_record_to_excel(rec1, self.test_excel)
        append_record_to_excel(rec2, self.test_excel)

        self.assertTrue(self.test_excel.exists())

        wb = openpyxl.load_workbook(self.test_excel)
        self.assertIn("Registros", wb.sheetnames)
        ws = wb["Registros"]

        # Row 1 is header, Row 2 is rec1, Row 3 is rec2
        self.assertEqual(ws.max_row, 3)
        self.assertEqual(ws.cell(row=2, column=6).value, "12345678")
        self.assertEqual(ws.cell(row=3, column=6).value, "009876543")
        self.assertEqual(ws.cell(row=3, column=7).value, "COLOMBIANA")


if __name__ == "__main__":
    unittest.main()
