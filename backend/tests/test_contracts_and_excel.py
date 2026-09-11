import unittest
import openpyxl
from openpyxl.comments import Comment
from pathlib import Path
from datetime import datetime
from backend.services.excel_analyzer import ExcelAnalyzer
from backend.services.worker_consolidator import WorkerConsolidator
from backend.services.contract_generator import ContractGenerator


class TestExcelAndContracts(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.test_dir = Path(__file__).parent / "test_data"
        cls.test_dir.mkdir(parents=True, exist_ok=True)
        cls.sample_excel = cls.test_dir / "Test_Planilla_Integral.xlsx"
        
        # Create rich test Excel with multiple sheets, hidden sheet, hidden row/col, comments
        wb = openpyxl.Workbook()
        
        # Sheet 1: DATOS (Visible)
        ws1 = wb.active
        ws1.title = "DATOS"
        ws1.append(["DNI", "Paterno", "Materno", "Nombres", "Cargo", "Area", "Remuneracion", "Direccion", "Celular"])
        ws1.append(["42694271", "GARCIA", "BRANDAN", "NADIA CELESTE", "OPERARIO DE PRODUCCION", "PLANTA", 1600.0, "JR. LAS FLORES 123", "999888777"])
        
        # Add comment with AFP note
        cell_remun = ws1.cell(row=2, column=7)
        cell_remun.comment = Comment("AFP HABITAT - Comision sobre Flujo", "RRHH")

        # Sheet 2: HISTORICO (Hidden Sheet)
        ws2 = wb.create_sheet("HISTORICO")
        ws2.sheet_state = "hidden"
        ws2.append(["Doc", "Apellidos", "Nombres", "Sueldo"])
        ws2.append(["10203040", "PEREZ LOPEZ", "JUAN CARLOS", 2500])

        # Sheet 3: AUXILIAR (With hidden row and hidden col)
        ws3 = wb.create_sheet("AUXILIAR")
        ws3.append(["DNI", "Estado Civil", "Asignacion", "AFP"])
        ws3.append(["42694271", "CASADO(A)", "SI", "HABITAT"])
        ws3.row_dimensions[2].hidden = True
        ws3.column_dimensions["C"].hidden = True

        wb.save(str(cls.sample_excel))

    def test_excel_analyzer_finds_worker_and_comments(self):
        analyzer = ExcelAnalyzer(self.sample_excel)
        result = analyzer.find_worker_by_dni("42694271")
        
        self.assertIsNotNone(result)
        self.assertEqual(result["dni"], "42694271")
        fields = result["fields"]
        
        self.assertIn("paternal_surname", fields)
        self.assertEqual(fields["paternal_surname"]["value"], "GARCIA")
        self.assertIn("afp_nota", fields)
        self.assertIn("AFP HABITAT", fields["afp_nota"]["value"])

    def test_excel_analyzer_finds_hidden_sheet(self):
        analyzer = ExcelAnalyzer(self.sample_excel)
        result = analyzer.find_worker_by_dni("10203040")
        
        self.assertIsNotNone(result)
        self.assertEqual(result["dni"], "10203040")
        self.assertIn("HISTORICO", result["sheets_found"])

    def test_contract_date_calculation_rule(self):
        # 08/09/2026 + 3 months -> 07/12/2026
        res = WorkerConsolidator.calculate_contract_dates("08/09/2026", 3)
        self.assertEqual(res["start_date"], "08/09/2026")
        self.assertEqual(res["end_date"], "07/12/2026")
        self.assertEqual(res["duration_text"], "3 meses")

        # 01/01/2026 + 1 month -> 31/01/2026
        res1 = WorkerConsolidator.calculate_contract_dates("01/01/2026", 1)
        self.assertEqual(res1["start_date"], "01/01/2026")
        self.assertEqual(res1["end_date"], "31/01/2026")

    def test_email_proposal_generation(self):
        email = WorkerConsolidator.generate_email_proposal("Nadia Celeste", "García Brandán", "empresa.com.pe")
        self.assertEqual(email, "nadia.garcia@empresa.com.pe")

    def test_number_to_soles_words(self):
        words = WorkerConsolidator.number_to_soles_words(1600.0)
        self.assertEqual(words, "MIL SEISCIENTOS Y 00/100 SOLES")
        
        words2 = WorkerConsolidator.number_to_soles_words("2,350.50")
        self.assertEqual(words2, "DOS MIL TRESCIENTOS CINCUENTA Y 50/100 SOLES")

    def test_contract_generation_preserves_template(self):
        consolidator = WorkerConsolidator(self.sample_excel)
        profile = consolidator.consolidate_worker_profile("42694271")
        
        dates = WorkerConsolidator.calculate_contract_dates("08/09/2026", 3)
        
        generator = ContractGenerator()
        gen_result = generator.generate_contract(profile, dates)
        
        self.assertEqual(gen_result["status"], "success")
        self.assertTrue(Path(gen_result["file_path"]).exists())
        self.assertIn("42694271", gen_result["file_name"])
        self.assertIn("GARCIA", gen_result["file_name"])


if __name__ == "__main__":
    unittest.main()
